"""Build reproducible Excel and HTML deliverables from the completed audit."""
import json, math, html, base64, hashlib, re
from pathlib import Path
from collections import Counter
import pandas as pd
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'FRAGMENT_DATA/recheck_20261019'
M=json.loads((OUT/'manifest.json').read_text())
DECISIONS={
4:('Not corroborated','No recurrent complete 62 bp haplotype. Apparent 62 bp alignment gaps are mostly portions of noisy larger-ITD reads.'),
6:('Off-size candidate','A distinct 27 bp haplotype is supported, but this does not establish the marked 28.78 bp fragment peak as a 29 bp ITD.'),
7:('Not corroborated','No recurrent complete 36/37 bp haplotype. Isolated alignment gaps near this size occur on larger/complex reads.'),
10:('Not corroborated','No coherent complete haplotype near 22 bp; a few inconsistent insertion alignments are insufficient.'),
14:('Not corroborated','No recurrent complete 40 bp haplotype; very few reads even have the corresponding net size.'),
15:('Not corroborated','No recurrent complete 46 bp haplotype. Nearby length-tail reads do not share a reproducible sequence.'),
17:('Not corroborated','Corrected swap: biological 11094 uses BAM 11202. No coherent 8 bp haplotype.'),
18:('Not corroborated','Corrected swap: biological 11094 uses BAM 11202. No coherent 11/12/13 bp haplotype.'),
20:('Not corroborated','Swap control: biological 11202 uses BAM 11094, which carries the expected 69 bp major ITD. No coherent minor ~21 bp haplotype.'),
24:('Not corroborated','No recurrent complete 30 bp haplotype; isolated shorter gaps arise in larger-ITD/error-containing reads.'),
33:('Sequence-supported','Coherent 21 bp insertion on both strands, confirmed competitively. The default GMM excludes almost all exact insertion-supporting reads.'),
38:('Not corroborated','No coherent 4 bp insertion haplotype. Short net gains form a broad sequencing-error background; lowering the size cutoff does not establish this peak.'),
42:('Weak off-size candidate','A low-count 27 bp haplotype exists, below the marked ~28.45 bp size and much lower than fragment AF. Retain as a review candidate.'),
43:('Not corroborated','No coherent 6 bp insertion haplotype; short length gains occur without a reproducible complete sequence.'),
46:('Not corroborated','Additional below-minimum check: no recurrent complete 11/12 bp haplotype.'),
49:('Not corroborated','No anchored read has net insertion size near 81 bp and no coherent 81 bp candidate was found.'),
55:('Sequence-supported','A 45 bp population occurs on both strands beside the known 48 bp ITD. Its two recurrent sequence forms differ by one junction base; this does not prove two biological ITD clones. In the baseline replay DADA2 separates a 227-read cluster, but its mixed 45 bp consensus has unresolved bases and is skipped.'),
}

def load(p): return pd.read_csv(p,sep='\t')
def calls_from_vcf(path):
    rows=[]
    for line in path.read_text().splitlines():
        if line.startswith('#') or not line: continue
        f=line.split('\t'); info=dict(x.split('=',1) for x in f[7].split(';') if '=' in x)
        rows.append(dict(vcf=str(path.relative_to(ROOT)),position=int(f[1]),insertion_length=int(info['ITD_LEN']),AF=float(info['AF']),DP=int(info['DP'])))
    return rows

summary=[]; mapping=[]; allele_rows=[]; profiles=[]; original_calls=[]; reruns=[]; links=[]
for sample in M['samples']:
    sid=sample['biological_sample']; p=OUT/sid
    stats=json.loads((p/'scan_complete.json').read_text()); astats=json.loads((p/'anchor_stats.json').read_text())
    anchors=load(p/'anchored_haplotypes.tsv'); sizes=load(p/'anchored_all_read_sizes.tsv.gz'); isolated=load(p/'allele_groups.tsv'); ins=load(p/'all_read_insertions.tsv.gz')
    mapping.append(dict(biological_sample=sid,physical_file_sample=sample['file_sample'],bam=sample['bam'],swap_applied=sid!=sample['file_sample'],trimmed_reads=stats['trimmed_min320'],reads_with_both_outer_anchors=astats['exact_both_anchor_reads'],default_GMM_assigned=sum(stats['default_gmm_counts'].values()),bam_read_groups=json.dumps(stats['bam_read_groups'])))
    prefix='clean_' if (p/'clean_validation_summary.tsv').exists() else 'anchor_'
    validation=load(p/(prefix+'validation_summary.tsv'))
    targets=json.loads((p/(prefix+'target_candidates.json')).read_text())
    for candidate in validation.to_dict('records'):
        target_alleles={x for values in targets.values() for x in values}
        known_lengths={int(r['Nanopore size (bp)']) for r in sample['known_calls']}
        role='target_candidate' if candidate['allele_id'] in target_alleles else 'known_ITD_sequence' if candidate['net_insertion_length'] in known_lengths else 'additional_diagnostic_haplotype_not_accepted_as_clone'
        candidate.update(biological_sample=sid,physical_file_sample=sample['file_sample'],validation_panel=prefix.rstrip('_'),candidate_role=role)
        allele_rows.append(candidate)
    for t in sample['targets']:
        row=t['excel_row']; fa=t['Fragment size (bp)']; decision,note=DECISIONS[row]
        ids=targets[str(row)]; selected=validation[validation.allele_id.isin(ids)]
        close=anchors[abs(anchors.net_insertion_length-fa)<=2.5]
        top=close.iloc[0] if len(close) else None
        exact=int(selected.n_reads.sum()) if len(selected) else int(top.n_reads) if top is not None else 0
        validated=int(selected.n_itd_reads.fillna(0).sum()) if len(selected) else None
        af=float(selected.allele_frequency.fillna(0).sum()) if len(selected) else None
        plus=int(selected.plus_reads.fillna(0).sum()) if len(selected) else None
        minus=int(selected.minus_reads.fillna(0).sum()) if len(selected) else None
        near=isolated[abs(isolated.insertion_length-fa)<=2.5].head(1)
        event_ids=set(ins[ins.allele.isin(near.allele)].read_id)
        actual_sizes=dict(Counter(sizes[sizes.read_id.isin(event_ids)].net_insertion_length))
        bestsizes=sorted(set(int(x) for x in selected.net_insertion_length)) if len(selected) else [int(top.net_insertion_length)] if top is not None else []
        summary.append(dict(excel_row=row,biological_sample=sid,physical_file_sample=sample['file_sample'],review_reason=t['review_reason'],fragment_size_bp=fa,fragment_AF=t['Fragment AF'],assessment=decision,candidate_net_sizes_bp=','.join(map(str,bestsizes)),exact_complete_haplotype_reads=exact,validated_candidate_reads=validated,validated_plus=plus,validated_minus=minus,candidate_AF=af,candidate_support_all_trimmed=float(selected.support_fraction_all_trimmed.fillna(0).sum()) if len(selected) else None,all_trimmed_reads=stats['trimmed_min320'],anchored_reads=astats['exact_both_anchor_reads'],reads_with_net_size_within_2_5bp=int(sum(abs(sizes.net_insertion_length-fa)<=2.5)),strongest_isolated_gap_bp=int(near.iloc[0].insertion_length) if len(near) else None,strongest_isolated_gap_reads=int(near.iloc[0].n_reads) if len(near) else 0,actual_net_sizes_on_these_reads=json.dumps(actual_sizes),conclusion=note,bam=sample['bam']))
        profiles.append((t,sizes))
    for mode in ('baseline','sensitive'):
        rp=p/(mode+'_rerun')
        if not (rp/'run_exit.json').exists(): continue
        exitinfo=json.loads((rp/'run_exit.json').read_text())
        vcfs=list((rp/'results').rglob('*.vcf'))
        calls=[r for v in vcfs for r in calls_from_vcf(v)]
        for call in calls or [{}]: reruns.append(dict(sample=sid,mode=mode,**exitinfo,**call))
        links.append(f'<li>{sid}: <a href="{sid}/{mode}_rerun/run.log">{mode} log</a>; <a href="{sid}/{mode}_rerun/command.json">exact command</a></li>')
    vcf_candidates=[f for f in (ROOT/'FRAGMENT_DATA/nanopore').glob(sample['file_sample']+'_*.vcf') if not f.name.startswith((sample['file_sample']+'_1_',sample['file_sample']+'_5_'))]
    preferred=[f for f in vcf_candidates if any(x in f.name for x in ('_2runs','_merged','_all_'))]
    for v in preferred or vcf_candidates:
        for call in calls_from_vcf(v): original_calls.append(dict(biological_sample=sid,physical_file_sample=sample['file_sample'],**call))

df=pd.DataFrame(summary).sort_values('excel_row'); ad=pd.DataFrame(allele_rows); md=pd.DataFrame(mapping); rd=pd.DataFrame(reruns)
df.to_csv(OUT/'recheck_summary.tsv',sep='\t',index=False)
ad.to_csv(OUT/'all_candidate_validation.tsv',sep='\t',index=False)
rd.to_csv(OUT/'parameter_rerun_calls.tsv',sep='\t',index=False)

notes=[
('Scope','14 YES-marked rows across 11 biological samples, plus two unmarked below-12-bp peaks and a swapped-sample control: 17 targets, 13 BAMs.'),
('Identity mapping','Excel biological 11094 -> physical BAM/report 11202; Excel biological 11202 -> physical BAM/report 11094. The 21/69-bp major calls verify the mapping; identity beyond the supplied swap annotation was not independently genotyped.'),
('Input integrity','COMPARISON.xlsx was read only. SHA256: '+M['workbook_sha256']),
('Runtime','Current workspace code mounted into local image teoloup/hematology-aml-flt3-itd:2.3; image ID sha256:c0f5c73116d5b67e9ea9056db8b6c7fc5e07880780a468b48d723b362a980ecd.'),
('FASTQ provenance','No original per-sample FASTQs were supplied in FRAGMENT_DATA. Region and primer-trimmed FASTQs were regenerated from each corresponding BAM and retained under <sample>/fastq. They are the same underlying sequencing evidence, not an independent assay.'),
('BAM selection','Use combined _2runs/_merged/_all where available; exclude _1/_5 dilution BAMs. Primary alignments and MAPQ >=20 were inspected for CIGAR evidence.'),
('Trimming','Cutadapt production primer settings, min read length 320 for discovery (default 330 replayed separately), maximum 666. All read IDs and qualities retained.'),
('Discovery','Compare a permissive gapped aligner, production aligner, original BAM CIGARs, and exact outer-anchor net-size/haplotype evidence. No GMM or ASV gate is imposed on all-read discovery.'),
('Outer anchors','WT[25:40] and WT[300:315], 0-based half-open. Window length in WT is 290 bp. Net insertion size is the complete observed window length minus 290, including all indels between anchors. Exact-anchor reads are a subset and their counts are conservative sequence evidence, not AF.'),
('Candidate panel','Complete exact anchored haplotypes with >=5 reads near target (+/-2.5 bp), up to three per target, plus known ITDs. Primary AF for the four candidate-bearing cases uses the clean panel. Extended diagnostic panels also contain recurrent neighboring error haplotypes; these are not accepted as biological ITDs.'),
('Validation','Current pipeline competitive alignment and breakpoint-gap checks on every trimmed read, with sequence caching expanded back to every ID and original strand. AF denominator is validated/classified reads; support fraction using all trimmed reads is also provided.'),
('Not corroborated','No reproducible full haplotype at the requested size in these data. This is not a calibrated negative test or proof of a fragment-analysis artifact. Tiny broad length tails and a few inconsistent reads are insufficient for confirmation.'),
('Small insertions','Exploratory discovery includes insertions >=3 bp and evaluates the below-minimum cases. None of the supplied below-12-bp targets is corroborated by a coherent full haplotype.'),
('Parameter experiments','Baseline replay: current default discovery with min reporting AF=0.005. Sensitive replay: min AF .001, min GMM fraction .001, max SD 10, subpeak fraction .01, min haplotype reads 10, max ITDs 10, min read length 320, min ITD 3, WT tolerance 1. These are review experiments, not a validated production configuration.'),
('Parameter results','11528: baseline reports 159,36,33 bp; sensitive settings add the supported 21-bp call at AF 2.303%, and additional low-level 24-bp/36-bp candidates. 13462: baseline reports only 48 bp; sensitive settings still skip the unresolved 45-bp population and add a 39-bp candidate at 0.187%. This demonstrates that broad parameter relaxation alone does not solve all missed calls; the additional candidates were not accepted automatically.'),
('Limitations','One supplied sequencing dataset per case; no UMIs or independent extraction/PCR replicates established. Distinct read IDs do not establish distinct original DNA molecules. Fragment electropherograms/raw peak traces and the original pipeline commands were unavailable.'),
('Shared allele','11528 21-bp allele matches the dominant 21-bp allele in 12156. This can reflect a recurrent hotspot allele; shared sequence alone cannot distinguish biological recurrence from cross-sample carryover.'),
('13462 failure trace','Baseline DADA2 infers two ASVs from the 503-read ITD peak. The 227-read ITD_1_H2 cluster has 170 extracted 45-bp insertions, including 103 of one sequence and 56 of another. Its context consensus is unresolved and skipped. The ambiguity summary then reports 0 Ns because it summarizes only surviving candidates.'),
('13462 sequence forms','The reconstructed 381-bp alleles differ at one base, T versus G at full-allele index 280. Their exact anchored supports are 57 reads (30+/27-) and 31 (15+/16-), with median base qualities 33 and 39 at that position. This supports the insertion-length population but does not establish whether the base difference is biological or technical.'),
('13462 read fate','Of the 211 reads assigned to the 45-bp candidate forms by the focused competitive panel, the baseline replay classifies 176 as validated support for its 48-bp reference. Another 28 prefer that reference but fail gap checks; seven prefer WT but fail checks. Candidate loss therefore also reallocates support to the surviving ITD.'),
('Small-gap example','In 11551 the production aligner finds 163 recurring GAGA insertions of 4 bp, but most linked reads have net gains of 12-18 bp, matching the 15-bp major ITD/error tail rather than a distinct 4-bp molecule. Recurrent alignment gaps alone can be misleading.'),
('Coordinate convention','Genomic anchor = hg38 amplicon start 28033881 + 0-based insertion boundary - 1. Repeat-equivalent insertion representations may have shifted anchor/insert strings.'),
]
improvements=[
('1','Discover candidates before GMM exclusion','11528: 226/228 reads with the recurring 21-bp insertion are excluded by the default GMM; two are assigned to WT. Length clustering should propose candidates, not limit the sequence search.'),
('2','Require consistent complete net length','Several apparent small gaps occur on larger-ITD molecules. Confirm both outer flanks and the whole local haplotype, rather than interpreting one of multiple alignment gaps as an independent ITD.'),
('3','Preserve sequence and breakpoint within each length','13462 has two recurrent 45-bp sequence forms, differing at one junction base, plus 48 bp. Length-only splitting cannot resolve this heterogeneity. Validate the supported full haplotypes and preserve uncertainty about the base difference rather than discard the shared insertion-length evidence.'),
('3b','Rescue and display unresolved candidate clusters','13462 already reaches a separate DADA2 cluster, but mixed 45-bp sequences yield an unresolved consensus that is skipped. Recluster the supported haplotypes and display rejected-candidate counts/ambiguity even when the remaining consensus has zero Ns.'),
('4','Separate discovery thresholds from reporting thresholds','A small clone beside a large peak can fail 15% of parent even with adequate total support. Use absolute read support and local error controls alongside global support; retain subthreshold candidates for review.'),
('5','Add explicit short-insertion review mode','For exploratory min size 3, WT tolerance must be <3 (e.g.1); search the WT peak and rescue before the no-non-WT early exit. Blindly lowering min size increases short error calls; require reproducible sequence and a local background comparison.'),
('6','Keep all eligible reads in final competition','Use WT + known ITDs + read-supported candidates; report ambiguous/rejected counts and all-trimmed support as well as the classified-read AF denominator.'),
('7','Do not treat every recurrent error haplotype as a clone','Neighboring -1/-2 bp versions of large ITDs can recur and be strand-biased. Use net size, complete sequence, quality/strand controls, independent PCR evidence, and background rather than automatic acceptance.'),
]

wb=openpyxl.Workbook(); wb.remove(wb.active)
frames={'Recheck_summary':df,'Sample_mapping':md,'Candidate_alleles':ad,'Original_calls':pd.DataFrame(original_calls),'Parameter_reruns':rd,'Methods':pd.DataFrame(notes,columns=['topic','detail']),'Improvements':pd.DataFrame(improvements,columns=['priority','change','evidence']),'Original_comparison':pd.DataFrame(M['rows'])}
for name,frame in frames.items():
    ws=wb.create_sheet(name)
    if frame.empty: ws.append(['No rows']); continue
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False,name=None):
        ws.append([None if isinstance(v,float) and (math.isnan(v) or math.isinf(v)) else str(v) if isinstance(v,(dict,list)) else v for v in row])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='183B56'); c.alignment=Alignment(wrap_text=True)
    for i,col in enumerate(frame.columns,1):
        width=55 if col in ('conclusion','detail','evidence','bam','actual_net_sizes_on_these_reads') else 22
        ws.column_dimensions[get_column_letter(i)].width=width
        if 'AF' in col or col=='candidate_support_all_trimmed':
            for cells in ws.iter_cols(min_col=i,max_col=i,min_row=2):
                for c in cells:
                    if isinstance(c.value,(int,float)): c.number_format='0.00%'
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    if name=='Recheck_summary':
        idx=list(frame.columns).index('assessment')+1
        for row in ws.iter_rows(min_row=2):
            value=row[idx-1].value; color='D7F0DF' if value=='Sequence-supported' else 'FFF0C2' if 'candidate' in value.lower() else 'EEF1F4'
            for c in row: c.fill=PatternFill('solid',fgColor=color)
xlsx=OUT/'FLT3_fragment_recheck.xlsx'; wb.save(xlsx)
reopened=openpyxl.load_workbook(xlsx,read_only=True,data_only=True)
assert reopened['Recheck_summary'].max_row==18
assert reopened['Sample_mapping'].max_row==14
assert hashlib.sha256((ROOT/'FRAGMENT_DATA/COMPARISON.xlsx').read_bytes()).hexdigest()==M['workbook_sha256']

flagged_profiles=[x for x in profiles if x[0]['review_reason']=='flagged']
fig,axes=plt.subplots(7,2,figsize=(13,19),constrained_layout=True)
for ax,(t,sizes) in zip(axes.flat,flagged_profiles):
    fa=t['Fragment size (bp)']; center=round(fa); lo=max(0,center-5); hi=center+5
    counts=sizes.net_insertion_length.value_counts(); xs=list(range(lo,hi+1))
    ax.bar(xs,[int(counts.get(x,0)) for x in xs],color='#577C94'); ax.axvline(fa,color='#CE6339',ls='--',lw=1.4)
    ax.set_title(f"{t['biological_sample']} | fragment {fa:.2f} bp | {DECISIONS[t['excel_row']][0]}",fontsize=10,loc='left')
    ax.set_xlabel('Net insertion size between outer anchors (bp)',fontsize=8); ax.set_ylabel('Reads',fontsize=8)
    ax.set_xticks(xs); ax.tick_params(labelsize=8); ax.spines[['top','right']].set_visible(False)
fig.suptitle('All-read net-size evidence at the 14 marked fragment peaks\nDashed line: fragment size. Bars include sequencing-error tails; sequence coherence is assessed separately.',fontsize=13)
fig.savefig(OUT/'marked_peak_net_sizes.png',dpi=140); plt.close(fig)

display=df[['biological_sample','physical_file_sample','fragment_size_bp','fragment_AF','assessment','candidate_net_sizes_bp','exact_complete_haplotype_reads','validated_candidate_reads','candidate_AF','conclusion']].copy()
display['fragment_size_bp']=display.fragment_size_bp.map(lambda x:f'{x:.2f}')
for col in ('fragment_AF','candidate_AF'): display[col]=display[col].map(lambda x:f'{100*x:.2f}%' if pd.notna(x) else '—')
display['validated_candidate_reads']=display.validated_candidate_reads.map(lambda x:str(int(x)) if pd.notna(x) else '—')
display=display.fillna('—')
table=display.to_html(index=False,escape=True,table_id='results',border=0)
article='https://pmc.ncbi.nlm.nih.gov/articles/PMC1907323/'
doc=f'''<!doctype html><html><head><meta charset="utf-8"><title>FLT3 fragment peak recheck</title><style>
body{{font:15px/1.5 system-ui,sans-serif;color:#203442;margin:32px auto;max-width:1500px;padding:0 24px}}h1,h2{{color:#183b56}}.lead{{font-size:19px;max-width:1100px}}.card{{background:#edf4f7;padding:18px 24px;border-radius:10px;margin:20px 0}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{text-align:left;vertical-align:top;padding:10px;border-bottom:1px solid #d9e3e8}}th{{background:#183b56;color:white;position:sticky;top:0}}tr:hover{{background:#f1f7fa}}input{{padding:10px;width:400px;max-width:90%;margin:10px 0}}.scroll{{overflow:auto}}img{{max-width:100%;height:auto}}code{{background:#f1f4f6;padding:2px 5px}}a{{color:#096987}}.small{{font-size:13px;color:#586a75}}
</style></head><body><h1>FLT3 fragment peak recheck</h1><p class="small">19 October 2026 · 17 reviewed targets · 13 BAMs · original workbook preserved</p>
<p class="lead">Sequence evidence supports the missed <b>11528 / 21 bp</b> population and <b>13462 / 45 bp</b> population. A separate 27 bp candidate in 10808 and a weak 27 bp candidate in 12156 differ from the marked fragment sizes. The remaining marked peaks lack a reproducible complete matching haplotype in these data.</p>
<div class="card"><b>Sample swap applied:</b> Excel biological <b>11094 → BAM 11202</b> (21 bp major ITD); Excel biological <b>11202 → BAM 11094</b> (69 bp major ITD). The names in the first table column follow Excel. <b>Evidence wording:</b> sequence-supported means corroborated in the supplied sequencing data, without establishing independent DNA molecules or a clinical limit of detection.</div>
<p><a href="FLT3_fragment_recheck.xlsx">Download the evidence workbook</a> · <a href="recheck_summary.tsv">Summary TSV</a> · <a href="manifest.json">Input mapping and workbook hash</a> · <a href="called_peak_plot_previews.html">Updated plots showing all final calls</a></p>
<h2>Per-peak results</h2><p>Candidate AF uses the pipeline's competitive classification/gap checks and a validated/classified-read denominator. Exact full-haplotype counts are conservative anchor evidence. A blank AF means no qualifying complete target candidate was proposed.</p><input id="filter" placeholder="Filter by sample, size, or assessment" oninput="for(const r of document.querySelectorAll('#results tbody tr'))r.hidden=!r.textContent.toLowerCase().includes(this.value.toLowerCase())"><div class="scroll">{table}</div>
<h2>Net-size profiles</h2><p>These measurements require both outer anchors. A small length tail alone does not demonstrate another ITD. Original BAM and permissive/production alignment tables are retained per sample.</p><img src="marked_peak_net_sizes.png" alt="Net insertion-size distributions at all fourteen marked fragment peaks">
<h2>Parameter experiments</h2><div class="scroll">{rd.to_html(index=False,border=0) if len(rd) else 'Pending'}</div><ul>{''.join(links)}</ul>
<h2>Algorithm changes supported by this audit</h2><ol>{''.join('<li><b>'+html.escape(title)+'.</b> '+html.escape(detail)+'</li>' for _,title,detail in improvements)}</ol>
<h2>Methods and practical limits</h2><dl>{''.join('<dt><b>'+html.escape(k)+'</b></dt><dd>'+html.escape(v)+'</dd>' for k,v in notes)}</dl>
<p>Low fragment peaks can require sequence confirmation; nonspecific or pull-up CE peaks are also possible. This is why absent corroboration is not labeled a proven artifact. <a href="{article}">Original FLT3 capillary-electrophoresis validation study</a>.</p>
</body></html>'''
(OUT/'FLT3_fragment_recheck.html').write_text(doc,encoding='utf-8')
print(df[['biological_sample','fragment_size_bp','assessment','candidate_net_sizes_bp','validated_candidate_reads','candidate_AF']].to_string(index=False))
print('Workbook reopened, dimensions validated, source workbook SHA256 unchanged.')
